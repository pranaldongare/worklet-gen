"""PDF and PPTX generators for Deep Research reports and Detailed Problem Statements."""

import io
import re
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from core.utils.generate_files import (
    normalize_text_list,
    format_multiline_pdf_bullet,
    ensure_list,
    estimate_height_wrapped_content,
    estimate_height_wrapped_Title,
    add_textbox,
    add_textbox_Title,
    CUSTOM_PAGE_SIZE,
    DEFAULT_PPT_GAP_INCH,
)

# ═══════════════════════════════════════════════════════════════════════
# PDF — Deep Research Report
# ═══════════════════════════════════════════════════════════════════════


def _pdf_styles():
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "dr_title", parent=styles["Heading1"], fontSize=22, textColor=colors.darkblue
    )
    section_style = ParagraphStyle(
        "dr_section",
        parent=styles["Heading2"],
        fontSize=16,
        textColor=colors.darkblue,
        spaceBefore=14,
        spaceAfter=6,
    )
    sub_header = ParagraphStyle(
        "dr_sub",
        parent=styles["Heading3"],
        fontSize=13,
        textColor=colors.HexColor("#333333"),
        spaceBefore=8,
        spaceAfter=4,
    )
    normal = ParagraphStyle("dr_normal", parent=styles["BodyText"], fontSize=11, leading=14)
    bullet = ParagraphStyle(
        "dr_bullet", parent=styles["BodyText"], fontSize=11, leftIndent=20, bulletIndent=10, leading=14
    )
    return title_style, section_style, sub_header, normal, bullet


def create_research_pdf(doc: dict, in_memory: bool = True) -> bytes | None:
    """Generate a PDF of the full deep research report."""
    try:
        title_style, section_style, sub_header, normal, bullet = _pdf_styles()
        elements = []

        # ── Title ──
        elements.append(Paragraph("Deep Research Report", title_style))
        elements.append(Spacer(1, 6))
        prompt = doc.get("prompt", "")
        if prompt:
            elements.append(Paragraph(f"<i>{escape(prompt[:500])}</i>", normal))
            elements.append(Spacer(1, 12))

        # ── Entities ──
        entities = doc.get("entities")
        if entities:
            elements.append(Paragraph("Entities Identified", section_style))
            for label, key in [
                ("Technologies", "technologies"),
                ("People", "people"),
                ("Companies", "companies"),
                ("Institutions", "institutions"),
                ("Research Areas", "research_areas"),
            ]:
                items = entities.get(key, [])
                if items:
                    elements.append(Paragraph(f"<b>{label}:</b>", normal))
                    for item in items:
                        elements.append(Paragraph(f"• {escape(str(item))}", bullet))
            elements.append(Spacer(1, 8))

        # ── As-Is ──
        as_is = doc.get("as_is")
        if as_is:
            elements.append(Paragraph("Current State of the Art", section_style))

            summary = as_is.get("summary", "")
            if summary:
                elements.append(Paragraph(escape(summary), normal))
                elements.append(Spacer(1, 6))

            sota_comparison = as_is.get("sota_comparison", []) or []
            if sota_comparison:
                elements.append(Paragraph("<b>Head-to-Head Comparison</b>", sub_header))
                for s in sota_comparison:
                    approach = s.get("approach", "")
                    actor = s.get("actor", "") or ""
                    metric = s.get("key_metric", "")
                    best = s.get("current_best", "")
                    strengths = s.get("strengths_one_line", "")
                    limitations = s.get("limitations_one_line", "")
                    source = s.get("source", "")
                    year = s.get("year", "")
                    actor_str = f" — {escape(str(actor))}" if actor else ""
                    year_str = f" ({year})" if year else ""
                    elements.append(
                        Paragraph(
                            f"<b>{escape(approach)}</b>{actor_str}{year_str}",
                            normal,
                        )
                    )
                    elements.append(
                        Paragraph(f"&nbsp;&nbsp;<b>{escape(metric)}:</b> {escape(best)}", bullet)
                    )
                    if strengths:
                        elements.append(Paragraph(f"&nbsp;&nbsp;+ {escape(strengths)}", bullet))
                    if limitations:
                        elements.append(Paragraph(f"&nbsp;&nbsp;- {escape(limitations)}", bullet))
                    if source:
                        if str(source).startswith("http"):
                            elements.append(
                                Paragraph(
                                    f'&nbsp;&nbsp;Source: <a href="{source}" color="blue"><u>link</u></a>',
                                    bullet,
                                )
                            )
                        else:
                            elements.append(
                                Paragraph(f"&nbsp;&nbsp;<i>Source: {escape(str(source))}</i>", bullet)
                            )
                elements.append(Spacer(1, 4))

            findings = as_is.get("key_findings", [])
            if findings:
                elements.append(Paragraph("<b>Key Findings</b>", sub_header))
                for f in findings:
                    finding = f.get("finding", "")
                    source = f.get("source", "")
                    year = f.get("year", "")
                    year_str = f" ({year})" if year else ""
                    source_str = f" — {escape(str(source))}" if source else ""
                    elements.append(
                        Paragraph(f"• {escape(finding)}{year_str}{source_str}", bullet)
                    )
                elements.append(Spacer(1, 4))

            timeline = as_is.get("timeline", [])
            if timeline:
                elements.append(Paragraph("<b>Timeline</b>", sub_header))
                for t in timeline:
                    year = t.get("year", "")
                    actor = t.get("actor", "")
                    event = t.get("event", "")
                    elements.append(
                        Paragraph(f"• <b>{year}</b>: {escape(actor)} — {escape(event)}", bullet)
                    )
                elements.append(Spacer(1, 4))

            players = as_is.get("key_players", [])
            if players:
                elements.append(Paragraph("<b>Key Players</b>", sub_header))
                for p in players:
                    name = p.get("name", "")
                    role = p.get("role", "")
                    work = p.get("notable_work", "")
                    link = p.get("link", "")
                    line = f"• <b>{escape(name)}</b> ({escape(role)}): {escape(work)}"
                    if link:
                        line += f' <a href="{link}" color="blue"><u>link</u></a>'
                    elements.append(Paragraph(line, bullet))
                elements.append(Spacer(1, 4))

        # ── Comparative ──
        comparative = doc.get("comparative")
        if comparative:
            elements.append(Paragraph("Comparative Analysis", section_style))

            comparisons = comparative.get("comparisons", [])
            if comparisons:
                elements.append(Paragraph("<b>Comparisons</b>", sub_header))
                for c in comparisons:
                    entity = c.get("entity", "")
                    approach = c.get("approach", "")
                    elements.append(
                        Paragraph(f"<b>{escape(entity)}</b>: {escape(approach)}", normal)
                    )
                    for s in c.get("strengths", []):
                        elements.append(Paragraph(f"&nbsp;&nbsp;+ {escape(s)}", bullet))
                    for l in c.get("limitations", []):
                        elements.append(Paragraph(f"&nbsp;&nbsp;- {escape(l)}", bullet))
                    elements.append(Spacer(1, 4))

            oss = comparative.get("open_source_landscape", [])
            if oss:
                elements.append(Paragraph("<b>Open Source Landscape</b>", sub_header))
                for p in oss:
                    name = p.get("name", "")
                    desc = p.get("description", "")
                    url = p.get("url", "")
                    stars = p.get("stars")
                    stars_str = f" ({stars} stars)" if stars else ""
                    line = f"• <b>{escape(name)}</b>{stars_str}: {escape(desc)}"
                    if url:
                        line += f' <a href="{url}" color="blue"><u>link</u></a>'
                    elements.append(Paragraph(line, bullet))
                elements.append(Spacer(1, 4))

            gaps = comparative.get("gaps", [])
            if gaps:
                elements.append(Paragraph("<b>Identified Gaps</b>", sub_header))
                for g in gaps:
                    elements.append(Paragraph(f"• {escape(str(g))}", bullet))
                elements.append(Spacer(1, 4))

        # ── Future Directions ──
        future = doc.get("future")
        if future:
            elements.append(Paragraph("Future Directions &amp; Problem Statements", section_style))

            problems = future.get("problem_statements", [])
            if problems:
                elements.append(Paragraph("<b>Forward-Looking Problem Statements</b>", sub_header))
                for p in problems:
                    title = p.get("title", "")
                    desc = p.get("description", "")
                    rationale = p.get("rationale", "")
                    impact = p.get("potential_impact", "")
                    core_tech = p.get("core_technologies", []) or []
                    research_areas = p.get("research_areas", []) or []
                    elements.append(Paragraph(f"<b>{escape(title)}</b>", normal))
                    if core_tech or research_areas:
                        chips = []
                        for t in core_tech:
                            chips.append(f"[Tech: {escape(str(t))}]")
                        for r in research_areas:
                            chips.append(f"[Area: {escape(str(r))}]")
                        elements.append(
                            Paragraph(f"<i>{' '.join(chips)}</i>", bullet)
                        )
                    elements.append(Paragraph(escape(desc), bullet))
                    if rationale:
                        elements.append(
                            Paragraph(f"<i>Rationale:</i> {escape(rationale)}", bullet)
                        )
                    if impact:
                        elements.append(
                            Paragraph(f"<i>Potential Impact:</i> {escape(impact)}", bullet)
                        )
                    elements.append(Spacer(1, 4))

            opps = future.get("opportunities", [])
            if opps:
                elements.append(Paragraph("<b>Opportunities</b>", sub_header))
                for o in opps:
                    elements.append(Paragraph(f"• {escape(str(o))}", bullet))
                elements.append(Spacer(1, 4))

            questions = future.get("research_questions", [])
            if questions:
                elements.append(Paragraph("<b>Open Research Questions</b>", sub_header))
                for i, q in enumerate(questions, 1):
                    if isinstance(q, dict):
                        question_text = q.get("question", "")
                        gain = q.get("expected_gain", "")
                        criteria = q.get("success_criteria", "")
                        elements.append(
                            Paragraph(f"<b>{i}. {escape(question_text)}</b>", normal)
                        )
                        if gain:
                            elements.append(
                                Paragraph(f"&nbsp;&nbsp;<i>Expected gain:</i> {escape(gain)}", bullet)
                            )
                        if criteria:
                            elements.append(
                                Paragraph(f"&nbsp;&nbsp;<i>Success criteria:</i> {escape(criteria)}", bullet)
                            )
                    else:
                        elements.append(Paragraph(f"{i}. {escape(str(q))}", bullet))
                elements.append(Spacer(1, 4))

        # ── References (grouped by source) ──
        refs = doc.get("references", [])
        if refs:
            elements.append(Paragraph("References", section_style))
            grouped: dict[str, list] = {}
            for ref in refs:
                if isinstance(ref, dict):
                    tag = (ref.get("tag") or "web").lower()
                else:
                    tag = "web"
                grouped.setdefault(tag, []).append(ref)

            group_order = ["patent", "scholar", "github", "web", "google"]
            tag_labels = {
                "patent": "Patents",
                "scholar": "Academic Papers",
                "github": "GitHub Repositories",
                "web": "Web References",
                "google": "Web References",
            }
            ordered_tags = [t for t in group_order if t in grouped] + [
                t for t in grouped if t not in group_order
            ]
            for tag in ordered_tags:
                group_refs = grouped[tag]
                label = tag_labels.get(tag, tag.title())
                elements.append(
                    Paragraph(f"<b>{label} ({len(group_refs)})</b>", sub_header)
                )
                for ref in group_refs:
                    if isinstance(ref, dict):
                        title = ref.get("title", "")
                        link = ref.get("link", "")
                        if title and link:
                            line = f'• {escape(title)} <a href="{link}" color="blue"><u>link</u></a>'
                        elif title:
                            line = f"• {escape(title)}"
                        elif link:
                            line = f'• <a href="{link}" color="blue"><u>link</u></a>'
                        else:
                            line = f"• {escape(str(ref))}"
                    else:
                        line = f"• {escape(str(ref))}"
                    elements.append(Paragraph(line, bullet))
                elements.append(Spacer(1, 4))

        if not elements:
            elements.append(Paragraph("No content available.", normal))

        if in_memory:
            buf = io.BytesIO()
            doc_pdf = SimpleDocTemplate(
                buf,
                pagesize=CUSTOM_PAGE_SIZE,
                leftMargin=40, rightMargin=40, topMargin=60, bottomMargin=40,
            )
            doc_pdf.build(elements)
            data = buf.getvalue()
            buf.close()
            return data
        return None
    except Exception as e:
        print(f"[generate_research_files] PDF creation failed: {e}")
        return None


# ═══════════════════════════════════════════════════════════════════════
# PPT — Deep Research Report
# ═══════════════════════════════════════════════════════════════════════


def _ppt_add_section_heading(slide, text, top_inch, gap):
    """Add a section heading to a PPT slide and return new top position."""
    left = Inches(0.5)
    top = Inches(top_inch)
    width = Inches(9.5)
    height = Inches(0.5)
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)
    return top_inch + 0.5 + gap


def create_research_ppt(doc: dict, in_memory: bool = True) -> bytes | None:
    """Generate a PPTX of the full deep research report."""
    try:
        prs = Presentation()
        prs.slide_width = Pt(750)
        prs.slide_height = Pt(1100)
        layout = prs.slide_layouts[6]

        slide = prs.slides.add_slide(layout)
        top = 0.5
        gap = DEFAULT_PPT_GAP_INCH
        top_margin = 0.5
        bottom_margin = 0.5
        slide_h = prs.slide_height.inches

        def _ensure_space(needed):
            nonlocal slide, top
            if top + needed > slide_h - bottom_margin:
                slide = prs.slides.add_slide(layout)
                top = top_margin

        # ── Title ──
        est = estimate_height_wrapped_Title("Deep Research Report")
        _ensure_space(est + gap)
        top = add_textbox_Title(slide, "Title", "Deep Research Report", top) + gap

        prompt = doc.get("prompt", "")
        if prompt:
            truncated = prompt[:300] + ("..." if len(prompt) > 300 else "")
            est = estimate_height_wrapped_content(truncated)
            _ensure_space(est + gap)
            top = add_textbox(slide, "Research Context", truncated, top)

        # ── Entities ──
        entities = doc.get("entities")
        if entities:
            _ensure_space(0.7)
            top = _ppt_add_section_heading(slide, "Entities Identified", top, gap)
            for label, key in [
                ("Technologies", "technologies"),
                ("People", "people"),
                ("Companies", "companies"),
                ("Institutions", "institutions"),
                ("Research Areas", "research_areas"),
            ]:
                items = entities.get(key, [])
                if items:
                    text = ", ".join(str(i) for i in items)
                    est = estimate_height_wrapped_content(text)
                    _ensure_space(est + gap)
                    top = add_textbox(slide, label, text, top)

        # ── As-Is ──
        as_is = doc.get("as_is")
        if as_is:
            _ensure_space(0.7)
            top = _ppt_add_section_heading(slide, "Current State of the Art", top, gap)

            summary = as_is.get("summary", "")
            if summary:
                est = estimate_height_wrapped_content(summary)
                _ensure_space(est + gap)
                top = add_textbox(slide, "Summary", summary, top)

            sota_comparison = as_is.get("sota_comparison", []) or []
            if sota_comparison:
                lines = []
                for s in sota_comparison:
                    approach = s.get("approach", "")
                    actor = s.get("actor", "") or ""
                    metric = s.get("key_metric", "")
                    best = s.get("current_best", "")
                    strengths = s.get("strengths_one_line", "")
                    limitations = s.get("limitations_one_line", "")
                    year = s.get("year", "")
                    actor_str = f" — {actor}" if actor else ""
                    year_str = f" ({year})" if year else ""
                    lines.append(f"• {approach}{actor_str}{year_str}")
                    lines.append(f"    {metric}: {best}")
                    if strengths:
                        lines.append(f"    + {strengths}")
                    if limitations:
                        lines.append(f"    - {limitations}")
                text = "\n".join(lines)
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, "Head-to-Head Comparison", text, top)

            findings = as_is.get("key_findings", [])
            if findings:
                lines = []
                for f in findings:
                    finding = f.get("finding", "")
                    year = f.get("year", "")
                    year_str = f" ({year})" if year else ""
                    lines.append(f"• {finding}{year_str}")
                text = "\n".join(lines)
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, "Key Findings", text, top)

            timeline = as_is.get("timeline", [])
            if timeline:
                lines = [f"• {t.get('year', '')}: {t.get('actor', '')} — {t.get('event', '')}" for t in timeline]
                text = "\n".join(lines)
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, "Timeline", text, top)

            players = as_is.get("key_players", [])
            if players:
                lines = [f"• {p.get('name', '')} ({p.get('role', '')}): {p.get('notable_work', '')}" for p in players]
                text = "\n".join(lines)
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, "Key Players", text, top)

        # ── Comparative ──
        comparative = doc.get("comparative")
        if comparative:
            _ensure_space(0.7)
            top = _ppt_add_section_heading(slide, "Comparative Analysis", top, gap)

            for c in comparative.get("comparisons", []):
                entity = c.get("entity", "")
                approach = c.get("approach", "")
                strengths = "\n".join(f"  + {s}" for s in c.get("strengths", []))
                limitations = "\n".join(f"  - {l}" for l in c.get("limitations", []))
                text = f"{approach}\n{strengths}\n{limitations}".strip()
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, entity, text, top)

            gaps = comparative.get("gaps", [])
            if gaps:
                text = "\n".join(f"• {g}" for g in gaps)
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, "Identified Gaps", text, top)

        # ── Future Directions ──
        future = doc.get("future")
        if future:
            _ensure_space(0.7)
            top = _ppt_add_section_heading(slide, "Future Directions", top, gap)

            for p in future.get("problem_statements", []):
                title = p.get("title", "")
                desc = p.get("description", "")
                rationale = p.get("rationale", "")
                impact = p.get("potential_impact", "")
                core_tech = p.get("core_technologies", []) or []
                research_areas = p.get("research_areas", []) or []
                chip_lines = []
                if core_tech:
                    chip_lines.append("Tech: " + ", ".join(str(t) for t in core_tech))
                if research_areas:
                    chip_lines.append("Areas: " + ", ".join(str(r) for r in research_areas))
                chips_str = ("\n" + "\n".join(chip_lines)) if chip_lines else ""
                text = f"{desc}{chips_str}\nRationale: {rationale}\nPotential Impact: {impact}"
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, title, text, top)

            opps = future.get("opportunities", [])
            if opps:
                text = "\n".join(f"• {o}" for o in opps)
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, "Opportunities", text, top)

            questions = future.get("research_questions", [])
            if questions:
                lines = []
                for i, q in enumerate(questions, 1):
                    if isinstance(q, dict):
                        question_text = q.get("question", "")
                        gain = q.get("expected_gain", "")
                        criteria = q.get("success_criteria", "")
                        lines.append(f"{i}. {question_text}")
                        if gain:
                            lines.append(f"    Expected gain: {gain}")
                        if criteria:
                            lines.append(f"    Success criteria: {criteria}")
                    else:
                        lines.append(f"{i}. {q}")
                text = "\n".join(lines)
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, "Open Research Questions", text, top)

        # ── References (grouped by source) ──
        refs = doc.get("references", [])
        if refs:
            _ensure_space(0.7)
            top = _ppt_add_section_heading(slide, "References", top, gap)

            grouped: dict[str, list] = {}
            for ref in refs:
                if isinstance(ref, dict):
                    tag = (ref.get("tag") or "web").lower()
                else:
                    tag = "web"
                grouped.setdefault(tag, []).append(ref)

            group_order = ["patent", "scholar", "github", "web", "google"]
            tag_labels = {
                "patent": "Patents",
                "scholar": "Academic Papers",
                "github": "GitHub Repositories",
                "web": "Web References",
                "google": "Web References",
            }
            ordered_tags = [t for t in group_order if t in grouped] + [
                t for t in grouped if t not in group_order
            ]
            for tag in ordered_tags:
                group_refs = grouped[tag]
                label = tag_labels.get(tag, tag.title())
                lines = []
                for ref in group_refs:
                    if isinstance(ref, dict):
                        title = ref.get("title", "")
                        link = ref.get("link", "")
                        lines.append(f"• {title}" + (f" — {link}" if link else ""))
                    else:
                        lines.append(f"• {ref}")
                text = "\n".join(lines)
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, f"{label} ({len(group_refs)})", text, top)

        if in_memory:
            buf = io.BytesIO()
            prs.save(buf)
            data = buf.getvalue()
            buf.close()
            return data
        return None
    except Exception as e:
        print(f"[generate_research_files] PPT creation failed: {e}")
        return None


# ═══════════════════════════════════════════════════════════════════════
# PDF — Detailed Problem Statement
# ═══════════════════════════════════════════════════════════════════════


def create_detailed_problem_pdf(problem: dict, in_memory: bool = True) -> bytes | None:
    """Generate a PDF of a single detailed problem statement."""
    try:
        title_style, section_style, sub_header, normal, bullet = _pdf_styles()
        elements = []

        title = problem.get("title", "Detailed Problem Statement")
        elements.append(Paragraph(escape(title), title_style))
        elements.append(Spacer(1, 10))

        # Text sections
        for label, key in [
            ("Executive Summary", "executive_summary"),
            ("Background &amp; Motivation", "background_and_motivation"),
            ("Problem Definition", "problem_definition"),
            ("Current State of the Art", "current_sota"),
            ("Proposed Approach", "proposed_approach"),
            ("Challenge / Use Case", "challenge_use_case"),
        ]:
            val = problem.get(key, "")
            if val:
                elements.append(Paragraph(f"<b>{label}</b>", section_style))
                elements.append(Paragraph(escape(str(val)), normal))
                elements.append(Spacer(1, 6))

        # List sections
        for label, key in [
            ("Deliverables", "deliverables"),
            ("KPIs", "kpis"),
            ("Prerequisites", "prerequisites"),
        ]:
            items = problem.get(key, [])
            if items:
                elements.append(Paragraph(f"<b>{label}</b>", section_style))
                for item in items:
                    text = format_multiline_pdf_bullet(str(item))
                    if text:
                        elements.append(Paragraph(text, bullet))
                elements.append(Spacer(1, 6))

        # Infra + Tech
        for label, key in [
            ("Infrastructure Requirements", "infrastructure_requirements"),
            ("Tech Stack", "tech_stack"),
        ]:
            val = problem.get(key, "")
            if val:
                elements.append(Paragraph(f"<b>{label}</b>", section_style))
                elements.append(Paragraph(escape(str(val)), normal))
                elements.append(Spacer(1, 6))

        # Milestones
        milestones = problem.get("milestones")
        if isinstance(milestones, dict) and milestones:
            elements.append(Paragraph("<b>Milestones (6 months)</b>", section_style))
            for key_ms in ("M2", "M4", "M6"):
                if key_ms in milestones:
                    elements.append(
                        Paragraph(f"• <b>{key_ms}:</b> {escape(str(milestones[key_ms]))}", bullet)
                    )
            for k, v in milestones.items():
                if k not in ("M2", "M4", "M6") and v:
                    elements.append(Paragraph(f"• <b>{k}:</b> {escape(str(v))}", bullet))
            elements.append(Spacer(1, 6))

        # Risk Assessment
        risks = problem.get("risk_assessment", [])
        if risks:
            elements.append(Paragraph("<b>Risk Assessment</b>", section_style))
            for r in risks:
                if isinstance(r, dict):
                    risk = r.get("risk", "")
                    impact = r.get("impact", "")
                    mitigation = r.get("mitigation", "")
                    elements.append(
                        Paragraph(
                            f"• <b>Risk:</b> {escape(risk)}<br/>"
                            f"&nbsp;&nbsp;<b>Impact:</b> {escape(impact)}<br/>"
                            f"&nbsp;&nbsp;<b>Mitigation:</b> {escape(mitigation)}",
                            bullet,
                        )
                    )
            elements.append(Spacer(1, 6))

        # Budget
        budget = problem.get("budget_estimation")
        if isinstance(budget, dict) and budget:
            elements.append(Paragraph("<b>Budget Estimation</b>", section_style))
            for k, v in budget.items():
                elements.append(Paragraph(f"• <b>{escape(str(k))}:</b> {escape(str(v))}", bullet))
            elements.append(Spacer(1, 6))

        # Key References
        key_refs = problem.get("key_references", [])
        if key_refs:
            elements.append(Paragraph("<b>Key References</b>", section_style))
            for ref in key_refs:
                elements.append(Paragraph(f"• {escape(str(ref))}", bullet))

        if in_memory:
            buf = io.BytesIO()
            doc_pdf = SimpleDocTemplate(
                buf,
                pagesize=CUSTOM_PAGE_SIZE,
                leftMargin=40, rightMargin=40, topMargin=60, bottomMargin=40,
            )
            doc_pdf.build(elements)
            data = buf.getvalue()
            buf.close()
            return data
        return None
    except Exception as e:
        print(f"[generate_research_files] Detailed problem PDF creation failed: {e}")
        return None


# ═══════════════════════════════════════════════════════════════════════
# PPT — Detailed Problem Statement
# ═══════════════════════════════════════════════════════════════════════


def create_detailed_problem_ppt(problem: dict, in_memory: bool = True) -> bytes | None:
    """Generate a PPTX of a single detailed problem statement."""
    try:
        prs = Presentation()
        prs.slide_width = Pt(750)
        prs.slide_height = Pt(1100)
        layout = prs.slide_layouts[6]

        slide = prs.slides.add_slide(layout)
        top = 0.5
        gap = DEFAULT_PPT_GAP_INCH
        top_margin = 0.5
        bottom_margin = 0.5
        slide_h = prs.slide_height.inches

        def _ensure_space(needed):
            nonlocal slide, top
            if top + needed > slide_h - bottom_margin:
                slide = prs.slides.add_slide(layout)
                top = top_margin

        # Title
        title = problem.get("title", "Detailed Problem Statement")
        est = estimate_height_wrapped_Title(title)
        _ensure_space(est + gap)
        top = add_textbox_Title(slide, "Title", title, top) + gap

        # Text sections
        for label, key in [
            ("Executive Summary", "executive_summary"),
            ("Background & Motivation", "background_and_motivation"),
            ("Problem Definition", "problem_definition"),
            ("Current State of the Art", "current_sota"),
            ("Proposed Approach", "proposed_approach"),
            ("Challenge / Use Case", "challenge_use_case"),
        ]:
            val = problem.get(key, "")
            if val:
                est = estimate_height_wrapped_content(str(val))
                _ensure_space(est + gap)
                top = add_textbox(slide, label, str(val), top)

        # List sections
        for label, key in [
            ("Deliverables", "deliverables"),
            ("KPIs", "kpis"),
            ("Prerequisites", "prerequisites"),
        ]:
            items = problem.get(key, [])
            if items:
                text = "\n".join(f"• {item}" for item in items)
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, label, text, top)

        # Infra + Tech
        for label, key in [
            ("Infrastructure Requirements", "infrastructure_requirements"),
            ("Tech Stack", "tech_stack"),
        ]:
            val = problem.get(key, "")
            if val:
                est = estimate_height_wrapped_content(str(val))
                _ensure_space(est + gap)
                top = add_textbox(slide, label, str(val), top)

        # Milestones
        milestones = problem.get("milestones")
        if isinstance(milestones, dict) and milestones:
            text = "\n".join(f"{k}: {v}" for k, v in milestones.items() if v)
            if text:
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, "Milestones (6 months)", text, top)

        # Risk Assessment
        risks = problem.get("risk_assessment", [])
        if risks:
            lines = []
            for r in risks:
                if isinstance(r, dict):
                    lines.append(
                        f"• Risk: {r.get('risk', '')}\n"
                        f"  Impact: {r.get('impact', '')}\n"
                        f"  Mitigation: {r.get('mitigation', '')}"
                    )
            text = "\n".join(lines)
            if text:
                est = estimate_height_wrapped_content(text)
                _ensure_space(est + gap)
                top = add_textbox(slide, "Risk Assessment", text, top)

        # Budget
        budget = problem.get("budget_estimation")
        if isinstance(budget, dict) and budget:
            text = "\n".join(f"• {k}: {v}" for k, v in budget.items())
            est = estimate_height_wrapped_content(text)
            _ensure_space(est + gap)
            top = add_textbox(slide, "Budget Estimation", text, top)

        # Key References
        key_refs = problem.get("key_references", [])
        if key_refs:
            text = "\n".join(f"• {ref}" for ref in key_refs)
            est = estimate_height_wrapped_content(text)
            _ensure_space(est + gap)
            top = add_textbox(slide, "Key References", text, top)

        if in_memory:
            buf = io.BytesIO()
            prs.save(buf)
            data = buf.getvalue()
            buf.close()
            return data
        return None
    except Exception as e:
        print(f"[generate_research_files] Detailed problem PPT creation failed: {e}")
        return None
